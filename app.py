import os
import tkinter as tk
from tkinter import filedialog, messagebox
import pyexcel_ods
import openpyxl
from odf.opendocument import load
from odf.text import P
from docx import Document
import traceback


# Function to convert a single .ods to .xlsx
def convert_ods_to_xlsx(input_path, output_path):
    try:
        print(f"Converting {input_path} to XLSX...")
        
        # Load .ods file using pyexcel-ods
        ods_data = pyexcel_ods.get_data(input_path)
        
        # Create a new workbook
        wb = openpyxl.Workbook()
        sheet = wb.active
        
        # Extract and write data from ods to the new .xlsx file
        for sheet_name, rows in ods_data.items():
            for row_index, row in enumerate(rows, start=1):
                for col_index, cell in enumerate(row, start=1):
                    sheet.cell(row=row_index, column=col_index, value=cell)
        
        # Save the new workbook
        output_file = os.path.join(output_path, os.path.splitext(os.path.basename(input_path))[0] + '.xlsx')
        wb.save(output_file)
        print(f"Successfully converted to {output_file}")
    except Exception as e:
        print("Error during ODS to XLSX conversion:")
        traceback.print_exc()
        raise Exception(f"Error converting {input_path}: {e}")


# Function to convert a single .odt to .docx
def convert_odt_to_docx(input_path, output_path):
    try:
        print(f"Converting {input_path} to DOCX...")
        
        # Load the .odt file
        doc = load(input_path)
        
        # Create a new .docx document
        docx_document = Document()

        # Extract paragraphs (text elements) from the .odt file
        paragraphs = doc.getElementsByType(P)

        for paragraph in paragraphs:
            text_content = ""
            for node in paragraph.childNodes:
                if node.nodeType == node.TEXT_NODE:
                    text_content += node.data

            if text_content.strip():
                docx_document.add_paragraph(text_content)

        # Save the document as .docx
        output_file = os.path.join(output_path, os.path.splitext(os.path.basename(input_path))[0] + '.docx')
        docx_document.save(output_file)
        
        print(f"Successfully converted to {output_file}")
    except Exception as e:
        print("Error during ODT to DOCX conversion:")
        traceback.print_exc()
        raise Exception(f"Error converting {input_path}: {e}")


# Bulk conversion handler
def bulk_convert(files, output_dir, conversion_type):
    for file in files:
        try:
            if conversion_type == "ods_to_xlsx":
                convert_ods_to_xlsx(file, output_dir)
            elif conversion_type == "odt_to_docx":
                convert_odt_to_docx(file, output_dir)
        except Exception as e:
            messagebox.showerror("Error", str(e))
            continue
    messagebox.showinfo("Success", f"Bulk conversion completed. Files saved in: {output_dir}")


# GUI Application
def main():
    def handle_convert():
        input_files = input_files_entry.get().split(";")
        output_dir = output_dir_entry.get()
        if not all(os.path.isfile(file) for file in input_files):
            messagebox.showerror("Error", "Please select valid input files.")
            return
        if not os.path.isdir(output_dir):
            messagebox.showerror("Error", "Please select a valid output directory.")
            return
        try:
            bulk_convert(input_files, output_dir, file_type_var.get())
        except Exception as e:
            messagebox.showerror("Error", f"Conversion failed: {e}")

    def select_files(file_type):
        file_types = [("ODS files", "*.ods")] if file_type == "ods_to_xlsx" else [("ODT files", "*.odt")]
        files = filedialog.askopenfilenames(filetypes=file_types)
        input_files_entry.delete(0, tk.END)
        input_files_entry.insert(0, ";".join(files))

    def select_folder():
        folder = filedialog.askdirectory()
        output_dir_entry.delete(0, tk.END)
        output_dir_entry.insert(0, folder)

    root = tk.Tk()
    root.title("File Converter")

    # File Type Selection
    tk.Label(root, text="Select Conversion Type:").grid(row=0, column=0, padx=10, pady=10, sticky="w")
    file_type_var = tk.StringVar(value="ods_to_xlsx")
    tk.Radiobutton(root, text="ODS to XLSX", variable=file_type_var, value="ods_to_xlsx").grid(row=0, column=1, padx=10, sticky="w")
    tk.Radiobutton(root, text="ODT to DOCX", variable=file_type_var, value="odt_to_docx").grid(row=0, column=2, padx=10, sticky="w")

    # Input Files
    tk.Label(root, text="Input Files:").grid(row=1, column=0, padx=10, pady=10, sticky="w")
    input_files_entry = tk.Entry(root, width=50)
    input_files_entry.grid(row=1, column=1, columnspan=2, padx=10, pady=10, sticky="w")
    tk.Button(root, text="Browse...", command=lambda: select_files(file_type_var.get())).grid(row=1, column=3, padx=10, pady=10)

    # Output Directory
    tk.Label(root, text="Output Directory:").grid(row=2, column=0, padx=10, pady=10, sticky="w")
    output_dir_entry = tk.Entry(root, width=50)
    output_dir_entry.grid(row=2, column=1, columnspan=2, padx=10, pady=10, sticky="w")
    tk.Button(root, text="Browse...", command=select_folder).grid(row=2, column=3, padx=10, pady=10)

    # Convert Button
    tk.Button(root, text="Convert", command=handle_convert, bg="green", fg="white").grid(row=3, column=1, columnspan=2, pady=20)

    root.mainloop()


if __name__ == "__main__":
    main()
